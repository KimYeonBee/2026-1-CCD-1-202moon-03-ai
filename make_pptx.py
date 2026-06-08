import pandas as pd
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

adhd = pd.read_excel('test.xlsx', sheet_name='ADHD').head(16)
non = pd.read_excel('test.xlsx', sheet_name='nonADHD').head(16)

conditions = ['A', 'B', 'C']
metrics = ['complete', 'quiz', 'quit', 'satisfaction']
metric_labels = ['완료율 (%)', '퀴즈 정답률 (%)', '이탈 횟수', '만족도']

C_ADHD = RGBColor(0x1A, 0x9A, 0xF5)
C_NON = RGBColor(0xBB, 0xBB, 0xBB)
C_DARK = RGBColor(0x22, 0x22, 0x22)
C_GRAY = RGBColor(0x88, 0x88, 0x88)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank

# 제목
tb = slide.shapes.add_textbox(Inches(0.5), Inches(0.25), Inches(12.3), Inches(0.8))
p = tb.text_frame.paragraphs[0]
r = p.add_run(); r.text = 'ADHD vs nonADHD  조건별 변화 추이'
r.font.size = Pt(28); r.font.bold = True; r.font.color.rgb = C_DARK
r.font.name = 'Apple SD Gothic Neo'

# 2x2 차트 배치
positions = [
    (Inches(0.4), Inches(1.15)),
    (Inches(6.85), Inches(1.15)),
    (Inches(0.4), Inches(4.3)),
    (Inches(6.85), Inches(4.3)),
]
cw, ch = Inches(6.1), Inches(3.0)

def style_series(series, color, width_pt=2.75):
    line = series.format.line
    line.color.rgb = color
    line.width = Pt(width_pt)
    # 마커
    from pptx.oxml.ns import qn
    sp = series._element
    # smooth off
    smooth = sp.find(qn('c:smooth'))
    if smooth is None:
        smooth = sp.makeelement(qn('c:smooth'), {})
        sp.append(smooth)
    smooth.set('val', '0')

for (metric, label), (x, y) in zip(zip(metrics, metric_labels), positions):
    cd = CategoryChartData()
    cd.categories = ['조건 A', '조건 B', '조건 C']
    adhd_vals = [round(float(adhd[f'{c}_{metric}'].mean()), 2) for c in conditions]
    non_vals = [round(float(non[f'{c}_{metric}'].mean()), 2) for c in conditions]
    cd.add_series('ADHD', adhd_vals)
    cd.add_series('nonADHD', non_vals)

    gf = slide.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, x, y, cw, ch, cd)
    chart = gf.chart
    chart.has_title = True
    chart.chart_title.text_frame.text = label
    tr = chart.chart_title.text_frame.paragraphs[0].runs[0]
    tr.font.size = Pt(15); tr.font.bold = True; tr.font.color.rgb = C_DARK
    tr.font.name = 'Apple SD Gothic Neo'

    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.TOP
    chart.legend.include_in_layout = False
    chart.legend.font.size = Pt(11); chart.legend.font.name = 'Apple SD Gothic Neo'

    s_adhd, s_non = chart.series[0], chart.series[1]
    style_series(s_adhd, C_ADHD, 3.0)
    style_series(s_non, C_NON, 2.75)

    # 데이터 라벨
    for s, col in [(s_adhd, C_ADHD), (s_non, C_GRAY)]:
        s.has_data_labels = True
        dl = s.data_labels
        dl.number_format = '0.0'; dl.number_format_is_linked = False
        dl.font.size = Pt(10); dl.font.bold = True; dl.font.color.rgb = col
        dl.font.name = 'Apple SD Gothic Neo'
        dl.position = XL_LABEL_POSITION.ABOVE

    # 축 폰트
    for ax in [chart.category_axis, chart.value_axis]:
        ax.tick_labels.font.size = Pt(10)
        ax.tick_labels.font.name = 'Apple SD Gothic Neo'

    # 값축 범위를 데이터에 맞춰 미리 설정 (차이가 잘 보이게, PPT에서 수정 가능)
    allv = adhd_vals + non_vals
    vmin, vmax = min(allv), max(allv)
    span = vmax - vmin
    base = span if span > 0 else max(vmax * 0.12, 0.6)
    va = chart.value_axis
    va.minimum_scale = round(max(0, vmin - base * 0.5), 2)
    va.maximum_scale = round(vmax + base * 0.4, 2)

prs.save('adhd_charts.pptx')
print('Saved adhd_charts.pptx')
